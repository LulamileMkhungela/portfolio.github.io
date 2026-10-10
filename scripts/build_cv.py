#!/usr/bin/env python3
"""Build the revised CV (PDF and DOCX) from Lulamile's 2026 CV.

Source of truth: Lulamile_Mkhungela_2026.pdf (root of the main branch).
This script keeps the facts from that CV and fixes wording only. Items that need a
decision are listed in docs/cv/CV-review-notes.md, not changed here.

Run from the repository root:
    python3 scripts/build_cv.py

Outputs:
    docs/cv/Lulamile-Mkhungela-CV.pdf
    docs/cv/Lulamile-Mkhungela-CV.docx

Requires: reportlab, python-docx.
"""
import html
from pathlib import Path

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer

OUT = Path(__file__).resolve().parent.parent / "docs" / "cv"

NAME = "LULAMILE MKHUNGELA"
TITLE = "Senior Product Designer, UI Developer & Frontend Developer"
CONTACT = [
    "mkhungela.l@gmail.com",
    "083 719 5064",
    "Johannesburg, South Africa",
    ("Portfolio", "https://lulamilemkhungela.github.io/portfolio.github.io/index.html"),
    ("LinkedIn", "https://www.linkedin.com/in/lulamile-mkhungela/"),
]
OPEN_TO = "Open to: Remote / Relocation / Freelance"
STRAPLINE = "UX/UI Design, Angular & React Development   |   8+ Years Full-Lifecycle Design & Development"

ACHIEVEMENTS = [
    "92% UAT acceptance on the Vodacom Engage employee app, the highest score in that digital portfolio at the time.",
    "Cut UI inconsistencies by 65% by standing up a new design system and Design Authority governance. HR-IT squads "
    "could review and request components, so products stayed consistent across Vodacom Opcos.",
    "Lifted Agile sprint velocity 35% by building a shared component library across two concurrent products (AddMoreDigital).",
    "Increased checkout completion 32% and average order value 18% by reducing a 7-step flow to 3 (RetailFlow).",
    "Shipped a 6-micro-frontend module-federation architecture and a 13-endpoint notifications system, in active "
    "production, for a national audit platform (AGSA).",
]

SUMMARY = [
    "Most product designers hand off to developers. I hand off to QA, because I have already built it. "
    "Senior Product Designer and Frontend Developer with 8+ years turning ambiguous product problems into shipped, "
    "measurable outcomes, not just interfaces.",
    "Consulting for the Auditor-General of South Africa, Eskom, Old Mutual, Sasol, Takeda, Toyota, and Vodacom as "
    "iOCO's embedded design-and-development authority on each account. I own the full software development lifecycle "
    "(SDLC), from user research through production React/Angular code, WCAG 2.2 AA compliance, and UAT. This is the "
    "'design engineer' capability enterprises are increasingly hiring for as the discipline consolidates from pure UI "
    "craft into blended product, data, and business-outcome ownership.",
    "Fluent on both sides of that blend: Figma and design systems on one end; Angular, TypeScript, module-federated "
    "micro-frontends, and AWS on the other. Extending into decision-support tooling, scenario/what-if analysis and "
    "experimentation, and Generative AI/RAG integration (Claude, Cursor, MCP Agents).",
]

SKILLS = [
    ("UX/UI & Product Design",
     "Figma, Figma Dev Mode, InVision, Framer, ProtoPie, Miro, Lovable, Adobe XD, Photoshop, Rapid Prototyping, Design "
     "Systems & Design Tokens, Information Architecture, UX Research, Usability Testing, Hotjar, Google Analytics, "
     "Journey Mapping, Service Blueprints, Interaction Design, Accessibility (WCAG 2.2 AA, ARIA), Screen Reader Testing"),
    ("Frontend Development",
     "React 19, Redux, State Management, Angular, Ionic, TypeScript, HTML5/CSS3, RxJS, NgRx, Angular Material, Angular CDK, "
     "Tailwind CSS, Bootstrap, SCSS, PWA, Vite, Leaflet, Storybook (Design-to-Dev Handoff), Component-Based Architecture, "
     "Responsive Design, Cross-Browser/Mobile-first Compatibility, Module Federation / Micro-Frontend Architecture, Vibe coding"),
    ("Full-Stack, Cloud & QA",
     "Node.js, RESTful API Design, AWS (S3, API Gateway, ECS Fargate), Firebase, Neo4j, OpenTelemetry, Android/Java, "
     "Database Design, UAT Support & Manual QA"),
    ("AI & Emerging Tech",
     "RAG (Retrieval-Augmented Generation), Multimodal AI, Claude, Claude Code, Cursor, MCP Agents, Generative AI "
     "Integration, v0"),
    ("Leadership & Collaboration",
     "Design Authority, Junior Mentoring, Stakeholder Interaction, Cross-Functional Delivery, Agile/Scrum, GitHub Actions, "
     "CI/CD Pipelines, Git, Pull Requests & Code Review, Jira, Azure DevOps"),
    ("Data & BI",
     "Highcharts, Recharts, Power BI, Power Pages, Canvas Apps, Decision-Support & Scenario-Analysis Dashboards"),
]

# Each role: (organisation, role line, dates, [bullets]). Sub-engagements under iOCO are nested.
IOCO_INTRO = ("Embedded as design-and-development authority across sequential and concurrent enterprise client "
              "engagements.")
IOCO_ENGAGEMENTS = [
    ("Eskom: Coal Stockpile Machine Learning Platform",
     "iOCO Consultant, embedded at Eskom. Frontend Developer & UX/UI Designer (React, Data Visualization)",
     "Aug 2026 – Present (Concurrent)",
     [
         "Own the dashboard and data-visualization layer for Eskom's coal-stockpile forecasting platform, turning multi-source "
         "ML outputs (coal supply, burn, quality, weather, and plant telemetry across multiple power stations) into interactive "
         "charts and drill-down reports for non-technical decision-makers.",
         "Design and build an interactive scenario-planning (\"what-if\") tool that lets Eskom managers vary input assumptions "
         "and see forecasted stock impact in real time, wired to prediction and risk-assessment endpoints from the ML backend.",
         "Partner with data engineering and data science teams to translate business and functional requirements (role-based "
         "access control, customizable stock-level widgets, audit-ready reporting) into an interface non-specialists can use.",
     ]),
    ("Auditor-General of South Africa (AGSA): Public Accountability Intelligence Platform",
     "iOCO Consultant, embedded at AGSA. Frontend Developer & UX/UI Designer (React, AWS)",
     "May 2026 – Present (Concurrent)",
     [
         "Build the public-facing frontend aggregating news, corruption data, audit findings, and statistical reporting, including "
         "a Leaflet/choropleth-based province risk heatmap, built on React 19 within a module-federation shell integrating 6 "
         "independently deployed micro-frontends.",
         "Design interactive data visualizations (trend graphs, statistical charts, comparative dashboards) through Figma "
         "prototyping and stakeholder review, then maintain WCAG 2.2 AA compliance across all public-facing surfaces as they go "
         "to production.",
         "Own the Media Risk Register and a full notifications system rebuild (13 REST endpoints, a tabbed alert center). Fixed a "
         "data-integrity bug where optimistic approve/reject updates were silently overwritten by auto-refresh cycles, restoring "
         "reliable audit-trail ordering for audit teams.",
     ]),
    ("Old Mutual: AI Governance Document Intelligence Platform",
     "iOCO Consultant, embedded at Old Mutual. Frontend Developer & Product Designer (React, AWS)",
     "May 2026 – Jun 2026",
     [
         "Designed and built the UI for an enterprise RAG platform that ingests and analyzes governance documents (PDF, image, "
         "video), including the DocumentUpload and Dashboard interfaces end-to-end, from Figma information architecture through "
         "production React/TypeScript. The aim was to replace manual search across thousands of documents with one verified, "
         "queryable knowledge base.",
         "Designed six regulatory compliance dashboards, wireframed and validated in Figma against compliance stakeholder review, "
         "mapped to FSCA/FAIS, SARB, FICA, POPIA, and Insurance Act requirements. Architected the AI findings pipeline so "
         "structured findings persisted to a database and surfaced dynamically at query time.",
         "Integrated React across AWS S3, API Gateway, Step Functions, and a Neo4j graph layer for storage, orchestration, and "
         "multimodal LLM retrieval spanning documents, images, and video transcripts. Carried the build through integration and "
         "QA testing in a regulated compliance environment before transitioning to the AGSA engagement.",
     ]),
    ("Takeda Pharmaceuticals: Global Supply Chain SaaS Dashboard",
     "iOCO Consultant, embedded at Takeda. UX/UI Design Specialist & UI Developer (Angular, Data Visualization)",
     "Apr 2025 – Apr 2026",
     [
         "Owned design and frontend delivery on a global pharmaceutical SaaS product reviewed by stakeholders across the US, "
         "India, and EMEA, working without a separate design-to-build handoff.",
         "Shipped a full Highcharts visualization suite (Heatmap, World Map, Timeline, Pareto, Mind Map), entirely "
         "settings-driven, and registered every component in Storybook for Takeda's shared Common Component Library, "
         "holding WCAG 2.2 AA across the entire product surface.",
         "Ran UAT and defect triage directly with stakeholders across three time zones, closing the design-to-QA loop.",
     ]),
    ("Toyota South Africa: Automotive Digital Platforms (Kinto, Automark, Remote App, Toyota App)",
     "iOCO Consultant, embedded at Toyota SA. Senior UI/UX Designer & Ionic/Angular Developer",
     "Aug 2023 – Apr 2025",
     [
         "Built a structured UX research practice using Hotjar across 10+ African markets, converting behavioral data into a "
         "prioritized design backlog tied directly to Toyota KPIs.",
         "Led the AltFind Settings redesign end-to-end, from research through Figma prototype to pixel-accurate production "
         "delivery, building the Ionic/Angular Material components myself.",
         "Delivered UI design in Figma and functional products (Ionic/Angular) for the Toyota Remote app's US-market release of "
         "new electric-vehicle features, and for Kinto's rollout of quick-approval flows for car rental and rent-to-own.",
         "Mentored four junior designers from research framing through component critique, and paired directly with the "
         "Angular engineering team on implementation reviews, closing gaps between Figma spec and shipped UI.",
     ]),
    ("Sasol: Enterprise Safety Reporting Platform (RCR)",
     "iOCO Consultant, embedded at Sasol. UI/UX Designer & Angular Developer",
     "Mar 2023 – Aug 2023",
     [
         "Designed and replaced a fragmented, Excel-based safety process with a live Angular platform adopted across five-plus "
         "departments, mapping a three-tier reporting hierarchy and NISS severity logic through stakeholder workshops.",
         "Delivered the full design process, design system, and end-to-end design, plus the Angular frontend with RxJS and "
         "RESTful API integration, eliminating manual data consolidation and giving leadership real-time visibility.",
         "Personally tested every release against NISS severity edge cases before rollout, and established UX governance "
         "standards that became the baseline for later Sasol digital safety tools.",
     ]),
    ("Vodacom: Enterprise HR & IT Digital Products",
     "iOCO Consultant, embedded at Vodacom. UI/UX Designer & Design System Lead",
     "Oct 2021 – Mar 2023",
     [
         "Designed the Vodacom Engage employee app in Figma, from first research session through delivery, passing testing at "
         "92% UAT acceptance (the highest score in the digital portfolio at the time) and launching across all Opcos.",
         "Architected the enterprise-wide Vodacom Design System in Figma, partnering directly with the Angular engineering team "
         "to validate components in code, as a single source of truth that accelerated engineering delivery.",
         "Cut UI inconsistencies by 65% through new Design Authority governance. Also delivered the Tobby AI chatbot interface, "
         "shipped with multilingual support across systems, and the External Bursary, Digital Visitor, and Engage App admin "
         "platforms, delivered and passed testing.",
     ]),
]

INDEPENDENT = [
    ("UluntuXd", "UX Facilitator & Coach", "Aug 2025 – Jul 2026 · Part-time, Remote",
     ["Designed and facilitated monthly MiCanvas workshops on the full product lifecycle for junior designers entering the "
      "South African product industry. Planned presentation design and ran design-critique sessions for a cohort of eight. "
      "Sent and received bi-weekly performance reports."]),
    ("IntellehubSA", "UX/UI Designer (DecisionGrid, DMS, PNS, Pillar, CAB Memo, Strategic Planning)",
     "Feb 2025 – Feb 2026 · Freelance",
     ["Applied a design-thinking process and designed six interactive prototypes for government platforms from scratch, "
      "in Figma for Power Pages and Canvas Apps. The designs follow South African government branding standards, and the "
      "internal team received a system they could maintain independently."]),
    ("AddMoreDigital", "UI/UX Designer, then Lead UI/UX Designer",
     "Sep 2021 – Sep 2022 · After-hours contract, concurrent with full-time role at iOCO",
     ["Designed and delivered Aziza, a GBV emergency-response app with real-time officer dispatch, a station dashboard, and "
      "live map tracking. Led concurrent Agile sprints across two products, managing five developers and three designers. "
      "A shared component library lifted sprint velocity by 35%."]),
    ("Hypothetical Objective Systems", "Senior UI/UX Designer & Developer Coach", "May 2020 – Dec 2021 · Freelance",
     ["Redesigned RetailFlow's checkout from seven steps to three, lifting checkout completion 32% and average order value "
      "18%. Restructured FarmTrack360's information architecture, cutting navigation time 47%."]),
    ("Sesyme & SmartServe", "UI/UX & Android Developer", "Jun 2019 – Mar 2020 · Contract concluded (startup funding lapsed)",
     ["Owned UI/UX design and native Android development end-to-end (wireframes, interactive design through to shipped "
      "solutions, user-tested features) for client-facing apps and websites (Sesyme app, SNB, SmartServe) on a small, "
      "fast-paced team."]),
    ("The Digital Academy", "Lead UI/UX Designer", "Aug 2018 – Jan 2019 · Fixed-term contract",
     ["Led UX/UI design for EntreHive's fintech and training platforms, translating stakeholder requirements into "
      "production-ready UI while mentoring junior Android developers. Actively assisted desktop teams with hands-on design "
      "work and guided their final design decisions through structured reviews."]),
    ("mLab", "UI/UX & Android Developer", "Apr 2017 – Feb 2018 · Fixed-term contract",
     ["Designed and built native Android interfaces for early-stage mobile products, including a co-working-space app, from "
      "user research through a working Java/Android application."]),
]

EDUCATION = [
    "Full-Stack Development Programme, FNB App Of The Year Academy, 2025 (Distinction, 92.7%)",
    "UX/UI Design Training, Wits JCSE, 2017",
    "NQF Level 5: Mobile & Web Development, MTN Business App Academy, 2021",
    "National Senior Certificate, Vaal Reefs Technical High School (2007–2011)",
    "Grade 12, Jeppe College of Commerce and Computer Studies (2016–2017)",
]
CERTIFICATIONS = (
    "Google UX Design Professional Certificate (Coursera); React.js (Pluralsight); Android Development Certificate; "
    "Angular: The Complete Guide; The Complete Claude Code & Claude Cowork Masterclass (Udemy); Software Development "
    "(LinkedIn Learning)"
)
AWARDS = "Winner, The Winning Team, Empire Foundation Hackathon"
REFERENCES = "References available on request."


def _esc(text: str) -> str:
    return html.escape(text, quote=False)


def build_pdf(path: Path) -> None:
    base = ParagraphStyle("base", fontName="Helvetica", fontSize=9, leading=12)
    small = ParagraphStyle("small", parent=base, fontSize=8.5, leading=11)
    name = ParagraphStyle("name", parent=base, fontName="Helvetica-Bold", fontSize=18, leading=22)
    title = ParagraphStyle("title", parent=base, fontName="Helvetica-Bold", fontSize=10.5, leading=14)
    h2 = ParagraphStyle("h2", parent=base, fontName="Helvetica-Bold", fontSize=10, leading=14,
                        spaceBefore=9, spaceAfter=3, textColor="#1f3b57")
    role = ParagraphStyle("role", parent=base, fontName="Helvetica-Bold", spaceBefore=5)
    sub = ParagraphStyle("sub", parent=base, fontName="Helvetica-Oblique", textColor="#333333")
    bullet = ParagraphStyle("bullet", parent=base, leftIndent=10, bulletIndent=0, spaceAfter=1)

    def bl(text):
        return Paragraph(_esc(text), bullet, bulletText="•")

    contact_bits = []
    for item in CONTACT:
        if isinstance(item, tuple):
            label, url = item
            contact_bits.append(f'<a href="{_esc(url)}" color="#1f3b57">{_esc(label)}</a>')
        else:
            contact_bits.append(_esc(item))
    story = [
        Paragraph(NAME, name),
        Paragraph(_esc(TITLE), title),
        Paragraph(" &nbsp;|&nbsp; ".join(contact_bits), small),
        Paragraph(_esc(OPEN_TO) + "<br/>" + _esc(STRAPLINE), small),
        Paragraph("KEY ACHIEVEMENTS", h2),
        *[bl(a) for a in ACHIEVEMENTS],
        Paragraph("PROFESSIONAL SUMMARY", h2),
        *[Paragraph(_esc(p), base) for p in SUMMARY],
        Paragraph("CORE SKILLS &amp; COMPETENCIES", h2),
        *[Paragraph(f"<b>{_esc(k)}:</b> {_esc(v)}", base) for k, v in SKILLS],
        Paragraph("PROFESSIONAL EXPERIENCE", h2),
        Paragraph("<b>iOCO, South Africa's Leading ICT Consulting Firm</b> · Oct 2021 – Present", role),
        Paragraph(_esc(IOCO_INTRO), sub),
    ]
    for client, role_line, dates, bullets in IOCO_ENGAGEMENTS:
        block = [
            Paragraph(f"<b>{_esc(client)}</b>", base),
            Paragraph(f"{_esc(role_line)} | {_esc(dates)}", sub),
            *[bl(b) for b in bullets],
            Spacer(1, 3),
        ]
        story.append(KeepTogether(block))
    story.append(Paragraph("INDEPENDENT &amp; CONTRACT EXPERIENCE", h2))
    for org, role_line, dates, bullets in INDEPENDENT:
        story.append(KeepTogether([
            Paragraph(f"<b>{_esc(org)}</b> · {_esc(role_line)}", base),
            Paragraph(_esc(dates), sub),
            *[bl(b) for b in bullets],
            Spacer(1, 3),
        ]))
    story += [
        Paragraph("EDUCATION, CERTIFICATIONS &amp; AWARDS", h2),
        *[bl(e) for e in EDUCATION],
        Paragraph(f"<b>Certifications:</b> {_esc(CERTIFICATIONS)}", base),
        Paragraph(f"<b>Award:</b> {_esc(AWARDS)}", base),
        Spacer(1, 6),
        Paragraph(_esc(REFERENCES), sub),
    ]
    doc = SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm,
        topMargin=14 * mm, bottomMargin=14 * mm, title="Lulamile Mkhungela CV", author="Lulamile Mkhungela",
    )
    doc.build(story)


def build_docx(path: Path) -> None:
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10)
    d.add_heading("Lulamile Mkhungela", level=1)
    d.add_paragraph(TITLE)
    d.add_paragraph(" | ".join(i[0] if isinstance(i, tuple) else i for i in CONTACT))
    d.add_paragraph(OPEN_TO)
    d.add_paragraph(STRAPLINE)
    d.add_heading("Key achievements", level=2)
    for a in ACHIEVEMENTS:
        d.add_paragraph(a, style="List Bullet")
    d.add_heading("Professional summary", level=2)
    for p in SUMMARY:
        d.add_paragraph(p)
    d.add_heading("Core skills and competencies", level=2)
    for k, v in SKILLS:
        p = d.add_paragraph()
        p.add_run(k + ": ").bold = True
        p.add_run(v)
    d.add_heading("Professional experience", level=2)
    p = d.add_paragraph()
    p.add_run("iOCO, South Africa's Leading ICT Consulting Firm · Oct 2021 – Present").bold = True
    d.add_paragraph(IOCO_INTRO)
    for client, role_line, dates, bullets in IOCO_ENGAGEMENTS:
        p = d.add_paragraph()
        p.add_run(client).bold = True
        d.add_paragraph(f"{role_line} | {dates}")
        for b in bullets:
            d.add_paragraph(b, style="List Bullet")
    d.add_heading("Independent and contract experience", level=2)
    for org, role_line, dates, bullets in INDEPENDENT:
        p = d.add_paragraph()
        p.add_run(f"{org} · {role_line}").bold = True
        d.add_paragraph(dates)
        for b in bullets:
            d.add_paragraph(b, style="List Bullet")
    d.add_heading("Education, certifications and awards", level=2)
    for e in EDUCATION:
        d.add_paragraph(e, style="List Bullet")
    d.add_paragraph("Certifications: " + CERTIFICATIONS)
    d.add_paragraph("Award: " + AWARDS)
    d.add_paragraph(REFERENCES)
    d.core_properties.author = "Lulamile Mkhungela"
    d.core_properties.title = "Lulamile Mkhungela CV"
    d.save(str(path))


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    build_pdf(OUT / "Lulamile-Mkhungela-CV.pdf")
    build_docx(OUT / "Lulamile-Mkhungela-CV.docx")
    print("wrote", OUT / "Lulamile-Mkhungela-CV.pdf", "and", OUT / "Lulamile-Mkhungela-CV.docx")
